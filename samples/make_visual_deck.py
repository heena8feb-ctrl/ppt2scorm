"""Creates a deck that exercises every kind of slide visual the extractor handles.
python samples/make_visual_deck.py samples/visuals_test.pptx"""
import io
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt


def png(w, h, draw, mode="RGB", bg="white"):
    im = Image.new(mode, (w, h), bg)
    draw(ImageDraw.Draw(im), w, h)
    b = io.BytesIO()
    im.save(b, "PNG")
    b.seek(0)
    return b


def photo(d, w, h):  # a big, noisy "photo" of a washing machine
    import random
    rnd = random.Random(4)
    for y in range(0, h, 4):
        for x in range(0, w, 4):
            v = 200 + rnd.randint(-30, 30)
            d.rectangle((x, y, x + 3, y + 3), fill=(v - 40, v - 20, v))
    d.rounded_rectangle((w * .3, h * .15, w * .7, h * .9), 30, fill=(240, 240, 245), outline=(60, 60, 70), width=8)
    d.ellipse((w * .38, h * .35, w * .62, h * .75), fill=(150, 170, 190), outline=(60, 60, 70), width=10)


def logo(d, w, h):
    d.rectangle((0, 0, w, h), fill=(20, 60, 120))
    d.text((10, h / 3), "ACME", fill="white")


def icon(d, w, h):
    d.ellipse((4, 4, w - 4, h - 4), fill=(240, 180, 0))


def filter_part(d, w, h):
    d.rounded_rectangle((60, 40, w - 60, h - 40), 40, outline=(30, 30, 30), width=10, fill=(230, 235, 240))
    for k in range(6):
        d.line((100 + k * 80, 70, 100 + k * 80, h - 70), fill=(80, 80, 80), width=6)


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    LOGO = png(300, 120, logo).getvalue()
    PART = png(900, 600, filter_part).getvalue()

    def slide(title, layout=5, notes=None):
        s = prs.slides.add_slide(prs.slide_layouts[layout])
        s.shapes.title.text = title
        s.shapes.add_picture(io.BytesIO(LOGO), Inches(11.4), Inches(6.6), width=Inches(1.6))  # logo on every slide
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    s = slide("Drain Filter Service: Visual Guide", 0)
    s.placeholders[1].text = "With photos, a chart and a process diagram"

    s = slide("Know the machine", notes="Point out where the drain filter cover is: bottom right on front loaders.")
    s.shapes.add_picture(png(2400, 1600, photo), Inches(1), Inches(1.5), width=Inches(7))       # big photo
    s.shapes.add_picture(png(64, 64, icon, "RGBA", (0, 0, 0, 0)), Inches(9), Inches(2), width=Inches(.5))  # icon
    tb = s.shapes.add_textbox(Inches(8.5), Inches(3), Inches(4), Inches(2)).text_frame
    tb.text = "The drain filter is behind the small cover at the bottom right."

    s = slide("The filter up close", notes="Lint, coins and hair pins are the usual culprits.")
    grp = s.shapes.add_group_shape()
    grp.shapes.add_picture(io.BytesIO(PART), Inches(1.5), Inches(1.6), width=Inches(6))        # picture in a group
    grp.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(6), Inches(.6)).text_frame.text = "Filter with debris trap"

    s = slide("Drain complaints by cause", notes="Half of all drain complaints are a blocked filter.")
    data = CategoryChartData()
    data.categories = ["Blocked filter", "Object in pump", "Kinked hose", "Other"]
    data.add_series("Share of complaints (%)", (50, 22, 18, 10))
    s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1.5), Inches(1.5), Inches(10), Inches(5), data)

    s = slide("The 4-step process", notes="Always unplug first, then drain, clean, refit and test.")
    steps = ["Unplug", "Drain", "Clean", "Refit & test"]
    boxes = []
    for k, label in enumerate(steps):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8 + k * 3.1), Inches(3), Inches(2.4), Inches(1.4))
        b.text_frame.text = label
        b.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
        boxes.append(b)
    for a, b in zip(boxes, boxes[1:]):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, a.left + a.width, a.top + a.height // 2, b.left, b.top + b.height // 2)
        c.line.width = Pt(3)

    s = slide("Recap: the filter", notes="Same close-up as before, reused.")
    s.shapes.add_picture(io.BytesIO(PART), Inches(3), Inches(1.6), width=Inches(5))             # same image again

    prs.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "visuals_test.pptx")
