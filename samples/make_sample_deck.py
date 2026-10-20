"""Creates a typical text-heavy corporate training deck to test the converter."""
from pptx import Presentation
from pptx.util import Inches, Pt

SLIDES = [
    ("Data Privacy Essentials", ["Mandatory training for all employees", "Duration: 15 minutes"],
     "Welcome to Data Privacy Essentials. This short course explains how we protect customer and employee data."),
    ("Why data privacy matters", [
        "Customers trust us with names, phone numbers, addresses and purchase history",
        "A single breach can lead to regulatory penalties and loss of trust",
        "India's DPDP Act 2023 places clear obligations on companies that process personal data",
        "Every employee is responsible, not just IT"],
     "Stress that privacy is everyone's job. Most breaches start with a human mistake, not a hacker."),
    ("What counts as personal data?", [
        "Any information that identifies a person directly or indirectly",
        "Examples: name, mobile number, email, address, Aadhaar, bank details",
        "Indirect identifiers: device ID, IP address, service ticket linked to a customer",
        "Sensitive data needs extra care: financial and health information"],
     "Ask learners to think of the customer data they touch daily, such as service tickets or warranty registrations."),
    ("The 5 golden rules", [
        "Collect only what you need",
        "Use data only for the purpose it was collected",
        "Share only with authorised people on approved tools",
        "Store it securely and never on personal devices or personal email",
        "Delete it when it is no longer needed"],
     "These five rules cover 90 percent of everyday situations."),
    ("Spotting a phishing attempt", [
        "Urgent tone: 'Your account will be blocked in 1 hour'",
        "Sender address that looks almost right",
        "Requests for passwords, OTPs or customer lists",
        "Unexpected attachments or links",
        "When in doubt, report it to the IT security mailbox, do not click"],
     "Phishing is the number one way attackers get in. Reporting a suspicious email is always the right move."),
    ("Real scenario: the urgent vendor request", [
        "A 'vendor' emails asking for last month's customer service call list",
        "The email says the CEO approved it and needs it in 30 minutes",
        "The sender domain is slightly different from the real vendor",
        "What should you do?"],
     "The right answer is to not send anything, verify through a known contact, and report the email. Urgency plus authority is a classic social engineering pattern."),
    ("If a breach happens", [
        "Report immediately to your manager and the IT security team",
        "Do not try to investigate or fix it yourself",
        "Do not delete evidence",
        "Speed matters: the law requires timely notification"],
     "Fast reporting limits damage. Nobody gets in trouble for reporting in good faith."),
    ("Key takeaways", [
        "Personal data is anything that identifies a person",
        "Follow the 5 golden rules every day",
        "Pause and verify before sharing data",
        "Report suspicious emails and incidents immediately"],
     "Thank you for completing the course."),
]


def main(out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, (title, bullets, notes) in enumerate(SLIDES):
        layout = prs.slide_layouts[0 if i == 0 else 1]
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            p = body.add_paragraph()
            p.text = b
            p.font.size = Pt(20)
        s.notes_slide.notes_text_frame.text = notes
    prs.save(out_path)


if __name__ == "__main__":
    import sys
    main(sys.argv[1] if len(sys.argv) > 1 else "sample_training.pptx")
