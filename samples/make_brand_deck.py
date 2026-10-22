"""Creates a deck with its own brand theme (colours + fonts) and a logo on every slide, to test brand detection.
python samples/make_brand_deck.py samples/branded_deck.pptx"""
import io
import re
import sys
import zipfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

BRAND = {"accent1": "0B6E4F", "accent2": "F2A541", "dk2": "1C2B36", "lt2": "F4F1EA"}
FONTS = {"majorFont": "Montserrat", "minorFont": "Georgia"}


def logo_png():
    im = Image.new("RGBA", (360, 120), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.ellipse((4, 10, 104, 110), fill=(11, 110, 79))
    d.rectangle((130, 40, 350, 80), fill=(28, 43, 54))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    logo = logo_png()
    slides = [("Greenfield Water Safety", None), ("Why clean water matters", ["Safe water prevents illness", "Filters need regular care"]),
              ("Changing a filter", ["Shut the valve", "Release pressure", "Swap the cartridge", "Flush for 5 minutes"]),
              ("Key takeaways", ["Change filters every 6 months", "Always flush after a change"])]
    for i, (title, bullets) in enumerate(slides):
        s = prs.slides.add_slide(prs.slide_layouts[0 if i == 0 else 1])
        s.shapes.title.text = title
        if bullets:
            tf = s.placeholders[1].text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                tf.add_paragraph().text = b
        s.shapes.add_picture(io.BytesIO(logo), Inches(11.6), Inches(6.7), width=Inches(1.4))
    buf = io.BytesIO()
    prs.save(buf)
    # rewrite the theme: brand colours and fonts
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"ppt/theme/theme\d+\.xml", item.filename):
                x = data.decode("utf-8")
                x = x.replace('name="Office Theme"', 'name="Greenfield"')
                for slot, val in BRAND.items():
                    x = re.sub(rf"(<a:{slot}>\s*<a:srgbClr val=\")[0-9A-Fa-f]{{6}}", rf"\g<1>{val}", x)
                for tag, fam in FONTS.items():
                    x = re.sub(rf"(<a:{tag}>\s*<a:latin typeface=\")[^\"]*", rf"\g<1>{fam}", x)
                data = x.encode("utf-8")
            zout.writestr(item, data)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "branded_deck.pptx")
