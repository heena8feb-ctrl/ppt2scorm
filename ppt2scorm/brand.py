"""Brand kits: logo, colours, fonts and corner style, applied to every course the player shows.

A kit is a folder:   <kit id>/kit.json   (+ logo.png, fonts/*.woff2|woff|ttf|otf)
kit.json:
  {"id", "name", "colors": {"primary", "accent", "text", "background"},
   "fonts": {"heading": FONT, "body": FONT}, "corners": "rounded" | "soft" | "square",
   "logo": "logo.png" | null, "logo_alt": str}
  FONT = {"stack": "<built-in id>"}  or  {"family": "Brand Sans", "files": [{"file", "weight", "style"}]}
         (a family without files is used only on computers that have it installed, with a safe fallback)

Kits can be detected from a deck's theme (colours, fonts, the logo repeated on slides), saved, exported
as one portable .json file (assets embedded) and imported elsewhere. Nothing loads from the internet:
fonts are bundled inside the SCORM package, so corporate firewalls and LMS restrictions can't break the look.
"""
import base64
import colorsys
import hashlib
import io
import json
import os
import re
import shutil
import zipfile
from pathlib import Path

HEX = re.compile(r"^#[0-9a-fA-F]{6}$")
MAX_KITS = int(os.environ.get("PPT2SCORM_MAX_BRAND_KITS", "50"))
SAFE = re.compile(r"^[A-Za-z0-9 _.&'()-]{1,60}$")
CORNERS = {"rounded": 1.0, "soft": 0.5, "square": 0.0}
FONT_EXTS = {".woff2": "woff2", ".woff": "woff", ".ttf": "truetype", ".otf": "opentype"}
MAX_LOGO = 4 * 1024 * 1024
MAX_FONT = 3 * 1024 * 1024
MAX_FONT_FILES = 6
SYSTEM = '-apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif'
# Fonts every learner already has (or a close match), so nothing needs downloading.
STACKS = {
    "system": ("System sans (Segoe UI, San Francisco, Roboto)", SYSTEM),
    "verdana": ("Verdana, wide and very legible", 'Verdana, "DejaVu Sans", Geneva, sans-serif'),
    "trebuchet": ("Trebuchet MS, friendly", '"Trebuchet MS", "Lucida Grande", "Segoe UI", sans-serif'),
    "arial": ("Arial / Helvetica", 'Arial, "Helvetica Neue", Helvetica, "Liberation Sans", sans-serif'),
    "calibri": ("Calibri, the PowerPoint default", 'Calibri, Carlito, "Segoe UI", Arial, sans-serif'),
    "georgia": ("Georgia, classic serif", 'Georgia, "Times New Roman", "DejaVu Serif", serif'),
    "cambria": ("Cambria, modern serif", 'Cambria, Caladea, Georgia, serif'),
}
DEFAULT = {
    "id": "standard", "name": "Standard look",
    "colors": {"primary": "#1f5eff", "accent": "#1f5eff", "text": "#17202e", "background": "#f3f5f9"},
    "fonts": {"heading": {"stack": "system"}, "body": {"stack": "system"}},
    "corners": "rounded", "logo": None, "logo_alt": "",
}
OFFICE_DEFAULT_ACCENTS = {"#4f81bd", "#4472c4", "#156082", "#5b9bd5", "#0f6fc6"}


class BrandError(ValueError):
    pass


# ---------- colour maths ----------
def _rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def _hex(rgb):
    return "#" + "".join(f"{max(0, min(255, round(c * 255))):02x}" for c in rgb)


def luminance(h):
    def ch(c):
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(c) for c in _rgb(h))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def mix(a, b, t):
    """t=0 -> a, t=1 -> b"""
    ra, rb = _rgb(a), _rgb(b)
    return _hex(tuple(x + (y - x) * t for x, y in zip(ra, rb)))


def readable_on(bg, dark="#111827"):
    """White or near-black, whichever reads better on bg."""
    return "#ffffff" if contrast(bg, "#ffffff") >= contrast(bg, dark) else dark


def darken_until(fg, bg, target=4.5):
    """Nudge fg darker (or lighter on dark bg) until it meets the contrast target against bg."""
    h, l, s = colorsys.rgb_to_hls(*_rgb(fg))
    step = -0.04 if luminance(bg) > 0.4 else 0.04
    for _ in range(30):
        c = _hex(colorsys.hls_to_rgb(h, l, s))
        if contrast(c, bg) >= target:
            return c
        l = max(0, min(1, l + step))
    return _hex(colorsys.hls_to_rgb(h, l, s))


def derived(kit):
    """All colours the player needs, derived so text always stays readable."""
    c = kit["colors"]
    bg, text, primary, accent = c["background"], c["text"], c["primary"], c["accent"]
    dark_page = luminance(bg) < 0.25
    card = mix(bg, "#ffffff", 0.08) if dark_page else "#ffffff"
    if contrast(text, card) < 4.5:
        text = readable_on(card)
    return {
        "brand": primary, "brand_ink": readable_on(primary), "accent": accent,
        "accent_text": darken_until(accent, card, 4.5),        # accent used as small text (labels) on cards
        "brand_text": darken_until(primary, card, 3.0),        # brand used for icons/borders on cards
        "ink": text, "muted": mix(text, card, 0.38), "bg": bg, "card": card,
        "line": mix(text, card, 0.86), "dark": dark_page,
    }


def check(kit):
    """Readability warnings in plain words (the kit still works; the player corrects what it must)."""
    c, d, out = kit["colors"], derived(kit), []
    if contrast(c["text"], d["card"]) < 4.5:
        out.append("The text colour is hard to read on the page; the course will use " +
                   ("white" if d["ink"] == "#ffffff" else "near-black") + " text instead.")
    if contrast(c["primary"], d["brand_ink"]) < 3:
        out.append("Buttons in the main colour will have low contrast. A darker or lighter main colour reads better.")
    if contrast(c["accent"], d["card"]) < 3:
        out.append("The accent colour is faint on white, so small labels use a darker shade of it.")
    return out


# ---------- validation ----------
def slug(name):
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40]
    return s or "brand"


def _font(v, where):
    if not isinstance(v, dict):
        raise BrandError(f"{where} font isn't valid")
    if "stack" in v:
        if v["stack"] not in STACKS:
            raise BrandError(f"{where} font: choose one from the list")
        return {"stack": v["stack"]}
    fam = v.get("family", "")
    if not isinstance(fam, str) or not SAFE.match(fam.strip()):
        raise BrandError(f"{where} font name can only use letters, numbers and spaces")
    files = []
    for f in v.get("files", [])[:MAX_FONT_FILES]:
        name = f.get("file", "") if isinstance(f, dict) else ""
        if not re.fullmatch(r"[a-z0-9_-]{1,60}\.(woff2|woff|ttf|otf)", name):
            raise BrandError(f"{where} font file isn't valid")
        weight = f.get("weight", 400)
        files.append({"file": name, "weight": weight if weight in (100, 200, 300, 400, 500, 600, 700, 800, 900) else 400,
                      "style": "italic" if f.get("style") == "italic" else "normal"})
    fallback = v.get("fallback", "system")
    return {"family": fam.strip(), "files": files, "fallback": fallback if fallback in STACKS else "system"}


def normalize(kit):
    """Validate a kit from a person or a file. Returns a clean copy; raises BrandError with a readable message."""
    if not isinstance(kit, dict):
        raise BrandError("the brand kit isn't valid")
    name = kit.get("name", "")
    if not isinstance(name, str) or not name.strip() or len(name) > 60:
        raise BrandError("give the brand kit a name (up to 60 characters)")
    colors = {}
    for k in ("primary", "accent", "text", "background"):
        v = (kit.get("colors") or {}).get(k, DEFAULT["colors"][k])
        if not isinstance(v, str) or not HEX.match(v):
            raise BrandError(f"{k} colour must look like #1f5eff")
        colors[k] = v.lower()
    fonts = {r: _font((kit.get("fonts") or {}).get(r, DEFAULT["fonts"][r]), r) for r in ("heading", "body")}
    corners = kit.get("corners", "rounded")
    if corners not in CORNERS:
        raise BrandError("corner style must be rounded, soft or square")
    logo = kit.get("logo")
    if logo is not None and logo != "logo.png":
        raise BrandError("the logo reference isn't valid")
    alt = kit.get("logo_alt", "")
    if not isinstance(alt, str) or len(alt) > 120:
        raise BrandError("logo description must be text under 120 characters")
    return {"id": slug(kit.get("id") or name), "name": name.strip(), "colors": colors, "fonts": fonts,
            "corners": corners, "logo": logo, "logo_alt": alt.strip() or name.strip()}


def clean_logo(data):
    """Re-encode an uploaded logo as PNG (max 800 px). Only real raster images get through; no SVG (it can hold script)."""
    if len(data) > MAX_LOGO:
        raise BrandError("the logo is over 4 MB")
    from PIL import Image
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
    except Exception:  # noqa: BLE001
        raise BrandError("the logo couldn't be read as an image. Use PNG, JPG or WebP.") from None
    if im.format not in ("PNG", "JPEG", "WEBP", "GIF"):
        raise BrandError("use a PNG, JPG or WebP logo")
    im = im.convert("RGBA")
    box = im.getbbox()                     # trim transparent margins
    if box:
        im = im.crop(box)
    im.thumbnail((800, 800))
    out = io.BytesIO()
    im.save(out, "PNG", optimize=True)
    return out.getvalue()


def check_font(data, filename):
    """Accept only genuine font files. Returns the extension to store it under."""
    ext = Path(filename or "").suffix.lower()
    if ext not in FONT_EXTS:
        raise BrandError("use a .woff2, .woff, .ttf or .otf font file")
    if len(data) > MAX_FONT:
        raise BrandError("that font file is over 3 MB")
    sig = data[:4]
    if ext == ".woff2":
        if sig != b"wOF2" or len(data) < 48 or int.from_bytes(data[8:12], "big") != len(data):
            raise BrandError("that file isn't a valid WOFF2 font")
        return ext
    if (ext == ".woff") != (sig == b"wOFF") or (ext != ".woff" and sig not in (b"\x00\x01\x00\x00", b"OTTO", b"true")):
        raise BrandError("that file isn't a valid font")
    try:
        from fontTools.ttLib import TTFont
        TTFont(io.BytesIO(data), lazy=True)["name"]
    except ImportError:
        pass
    except Exception:  # noqa: BLE001
        raise BrandError("that font file is damaged or isn't a font") from None
    return ext


# ---------- detection from a deck ----------
def from_deck(pptx_path):
    """Build a kit from the deck's theme and the logo repeated on its slides.
    Returns (kit, logo_png_bytes or None, notes)."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from .extract import DeckError, check_archive

    try:
        check_archive(pptx_path)
    except DeckError as e:
        raise BrandError(str(e)) from None
    notes = []
    theme = {}
    with zipfile.ZipFile(pptx_path) as z:
        names = sorted(n for n in z.namelist() if re.fullmatch(r"ppt/theme/theme\d+\.xml", n))
        xml = z.read(names[0]).decode("utf-8", "replace") if names else ""
    for slot in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2"):
        m = re.search(rf"<a:{slot}>\s*<a:(?:srgbClr val|sysClr val=\"[^\"]*\" lastClr)=\"([0-9A-Fa-f]{{6}})\"", xml)
        if m:
            theme[slot] = "#" + m.group(1).lower()
    fonts = {}
    for role, tag in (("heading", "majorFont"), ("body", "minorFont")):
        m = re.search(rf"<a:{tag}>\s*<a:latin typeface=\"([^\"]*)\"", xml)
        if m and m.group(1).strip():
            fonts[role] = m.group(1).strip()
    tname = (re.search(r'<a:theme [^>]*name="([^"]*)"', xml) or [None, ""])[1]

    primary = theme.get("accent1", DEFAULT["colors"]["primary"])
    accent = theme.get("accent2", primary)
    text = theme.get("dk2") if theme.get("dk1") in (None, "#000000") and theme.get("dk2") and luminance(theme["dk2"]) < 0.12 \
        else theme.get("dk1", DEFAULT["colors"]["text"])
    if text == "#000000":
        text = "#161c24"
    bg = mix(theme.get("lt1", "#ffffff"), primary, 0.05)
    if primary in OFFICE_DEFAULT_ACCENTS or tname in ("Office Theme", "Office"):
        notes.append("This deck uses PowerPoint's built-in Office theme, so its colours and fonts may not be your brand's.")

    def font(fam):
        key = {"calibri": "calibri", "arial": "arial", "verdana": "verdana", "georgia": "georgia", "cambria": "cambria",
               "trebuchet ms": "trebuchet", "segoe ui": "system"}.get(fam.lower())
        if key:
            return {"stack": key}
        if not SAFE.match(fam):
            return {"stack": "system"}
        notes.append(f"Upload the font files for \"{fam}\" so it shows on every learner's computer; until then it's "
                     "used only where it's installed.")
        return {"family": fam, "files": [], "fallback": "system"}

    kit_fonts = {r: font(fonts[r]) if r in fonts else {"stack": "system"} for r in ("heading", "body")}
    if kit_fonts["heading"] == kit_fonts["body"] and len([n for n in notes if "Upload the font" in n]) > 1:
        notes.pop()

    # logo: a picture on the slide master or layouts, else one repeated on many slides; never a full-bleed background
    prs = Presentation(pptx_path)
    area = prs.slide_width * prs.slide_height
    cands = {}

    def consider(shapes, weight):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                consider(sh.shapes, weight)
                continue
            if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                img = sh.image
            except Exception:  # noqa: BLE001
                continue
            if (sh.width or 0) * (sh.height or 0) > 0.25 * area:
                continue
            c = cands.setdefault(img.sha1, {"blob": img.blob, "score": 0})
            c["score"] += weight

    consider(prs.slide_master.shapes, 5)
    for lay in prs.slide_layouts:
        consider(lay.shapes, 2)
    n = len(prs.slides)
    per_slide = {}
    for s in prs.slides:
        seen = set()
        for sh in s.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    seen.add(sh.image.sha1)
                except Exception:  # noqa: BLE001
                    pass
        for h in seen:
            per_slide[h] = per_slide.get(h, 0) + 1
    for s in prs.slides:
        for sh in s.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    h = sh.image.sha1
                except Exception:  # noqa: BLE001
                    continue
                if per_slide[h] >= max(3, (n + 1) // 2) and h not in cands and (sh.width or 0) * (sh.height or 0) <= 0.25 * area:
                    cands[h] = {"blob": sh.image.blob, "score": per_slide[h]}
    logo = None
    if cands:
        best = max(cands.values(), key=lambda c: c["score"])
        try:
            logo = clean_logo(best["blob"])
        except BrandError:
            logo = None
    if not logo:
        notes.append("No logo was found on the slide master or repeated across slides. Add one in the brand kit.")

    title = (Path(pptx_path).stem if isinstance(pptx_path, (str, Path)) else "Deck").replace("_", " ").strip()[:40] or "Deck"
    kit = normalize({"name": f"{title} brand", "colors": {"primary": primary, "accent": accent, "text": text, "background": bg},
                     "fonts": kit_fonts, "corners": "rounded", "logo": "logo.png" if logo else None, "logo_alt": ""})
    return kit, logo, notes


# ---------- storage ----------
class Store:
    """Kits on disk: <root>/<id>/kit.json + assets."""

    def __init__(self, root):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)

    def path(self, kid):
        if not re.fullmatch(r"[a-z0-9-]{1,40}", kid or ""):
            raise BrandError("unknown brand kit")
        return self.root / kid

    def list(self):
        out = []
        for d in sorted(self.root.iterdir()):
            f = d / "kit.json"
            if f.is_file():
                try:
                    out.append(json.loads(f.read_text(encoding="utf-8")))
                except ValueError:
                    continue
        return sorted(out, key=lambda k: k["name"].lower())

    def get(self, kid):
        f = self.path(kid) / "kit.json"
        if not f.is_file():
            raise BrandError("that brand kit no longer exists")
        return json.loads(f.read_text(encoding="utf-8"))

    def save(self, kit, logo=None, fonts=None, replace_id=None):
        """kit: dict from a person (validated here). logo: PNG bytes or None (keep existing).
        fonts: {filename: bytes} to add. Returns the saved kit."""
        kit = normalize(kit)
        if replace_id:
            kit["id"] = replace_id
        else:
            if self.root.is_dir() and sum(1 for d in self.root.iterdir() if d.is_dir()) >= MAX_KITS:
                raise BrandError(f"this server already has {MAX_KITS} brand kits. Delete one you no longer use first")
            base, i = kit["id"], 2
            while (self.root / kit["id"]).exists():
                kit["id"] = f"{base}-{i}"[:40]
                i += 1
        d = self.path(kit["id"])
        (d / "fonts").mkdir(parents=True, exist_ok=True)
        if logo is not None:
            (d / "logo.png").write_bytes(logo)
        kit["logo"] = "logo.png" if (d / "logo.png").is_file() and kit.get("logo") is not None or logo else None
        for name, data in (fonts or {}).items():
            (d / "fonts" / name).write_bytes(data)
        used = {f["file"] for r in ("heading", "body") for f in kit["fonts"][r].get("files", [])}
        for r in ("heading", "body"):
            kit["fonts"][r]["files"] = [f for f in kit["fonts"][r].get("files", []) if (d / "fonts" / f["file"]).is_file()] \
                if "files" in kit["fonts"][r] else kit["fonts"][r].get("files")
            if kit["fonts"][r].get("files") is None:
                kit["fonts"][r].pop("files", None)
        for f in (d / "fonts").iterdir():   # drop font files no longer referenced
            if f.name not in used:
                f.unlink()
        (d / "kit.json").write_text(json.dumps(kit, indent=2, ensure_ascii=False), encoding="utf-8")
        return kit

    def delete(self, kid):
        shutil.rmtree(self.path(kid), ignore_errors=True)

    def export(self, kid):
        return export_kit(self.path(kid))

    def import_(self, data):
        kit, logo, fonts = parse_export(data)
        return self.save(kit, logo=logo, fonts=fonts)


def export_kit(kit_dir):
    """One portable JSON file with the logo and fonts embedded."""
    kit_dir = Path(kit_dir)
    kit = json.loads((kit_dir / "kit.json").read_text(encoding="utf-8"))
    assets = {}
    if (kit_dir / "logo.png").is_file():
        assets["logo.png"] = base64.b64encode((kit_dir / "logo.png").read_bytes()).decode()
    for f in sorted((kit_dir / "fonts").glob("*")) if (kit_dir / "fonts").is_dir() else []:
        assets["fonts/" + f.name] = base64.b64encode(f.read_bytes()).decode()
    return {"ppt2scorm_brand_kit": 1, "kit": kit, "assets": assets}


def parse_export(data):
    try:
        doc = json.loads(data) if isinstance(data, (bytes, str)) else data
    except ValueError:
        raise BrandError("that file isn't a ppt2scorm brand kit") from None
    if not isinstance(doc, dict) or doc.get("ppt2scorm_brand_kit") != 1 or not isinstance(doc.get("assets", {}), dict):
        raise BrandError("that file isn't a ppt2scorm brand kit")
    kit = normalize(doc.get("kit"))
    assets = doc.get("assets") or {}
    logo = None
    if "logo.png" in assets:
        logo = clean_logo(base64.b64decode(assets["logo.png"]))
    fonts = {}
    for key, b64 in assets.items():
        if key.startswith("fonts/"):
            name = key[6:]
            if not re.fullmatch(r"[a-z0-9_-]{1,60}\.(woff2|woff|ttf|otf)", name):
                raise BrandError("the kit file contains an unexpected font name")
            raw = base64.b64decode(b64)
            check_font(raw, name)
            fonts[name] = raw
    return kit, logo, fonts


def font_file_name(family, data, ext):
    return f"{slug(family)}-{hashlib.sha1(data).hexdigest()[:8]}{ext}"


def write_kit_dir(kit_dir, kit, logo=None, fonts=None):
    """Write a kit (already normalized) into a folder, e.g. a course's own brand snapshot."""
    kit_dir = Path(kit_dir)
    if kit_dir.exists():
        shutil.rmtree(kit_dir)
    (kit_dir / "fonts").mkdir(parents=True)
    if logo:
        (kit_dir / "logo.png").write_bytes(logo)
    for name, data in (fonts or {}).items():
        (kit_dir / "fonts" / name).write_bytes(data)
    kit = dict(kit, logo="logo.png" if logo or (kit_dir / "logo.png").is_file() else None)
    (kit_dir / "kit.json").write_text(json.dumps(kit, indent=2, ensure_ascii=False), encoding="utf-8")
    return kit


def copy_kit_dir(src, dst):
    if Path(dst).exists():
        shutil.rmtree(dst)
    shutil.copytree(src, dst)


# ---------- output for the player ----------
def _stack(f):
    if "stack" in f:
        return STACKS[f["stack"]][1]
    return f'"{f["family"]}", ' + STACKS[f.get("fallback", "system")][1]


def css(kit, font_url="brand/fonts/"):
    """brand.css for the player: @font-face rules plus colour, font and corner tokens."""
    d = derived(kit)
    out = ["/* Brand kit: " + kit["name"].replace("*/", "") + " */"]
    seen = set()
    for r in ("heading", "body"):
        f = kit["fonts"][r]
        for ff in f.get("files", []):
            if ff["file"] in seen:
                continue
            seen.add(ff["file"])
            fmt = FONT_EXTS[Path(ff["file"]).suffix]
            out.append(f'@font-face {{ font-family: "{f["family"]}"; src: url("{font_url}{ff["file"]}") format("{fmt}"); '
                       f'font-weight: {ff["weight"]}; font-style: {ff["style"]}; font-display: swap; }}')
    out.append(":root {")
    for k, v in (("brand", d["brand"]), ("brand-ink", d["brand_ink"]), ("accent", d["accent"]),
                 ("accent-text", d["accent_text"]), ("brand-text", d["brand_text"]), ("ink", d["ink"]),
                 ("muted", d["muted"]), ("bg", d["bg"]), ("card", d["card"]), ("line", d["line"])):
        out.append(f"  --{k}: {v};")
    out.append(f"  --font-body: {_stack(kit['fonts']['body'])};")
    out.append(f"  --font-heading: {_stack(kit['fonts']['heading'])};")
    out.append(f"  --rs: {CORNERS[kit['corners']]:g};")
    out.append("}")
    return "\n".join(out) + "\n"


def package_into(build_dir, kit_dir):
    """Copy a course's brand snapshot into a package folder and write brand.css. Returns theme info for course.js."""
    build_dir, kit_dir = Path(build_dir), Path(kit_dir)
    kit = normalize(json.loads((kit_dir / "kit.json").read_text(encoding="utf-8")))
    (build_dir / "brand").mkdir(exist_ok=True)
    theme = {"brand": kit["colors"]["primary"], "name": kit["name"]}
    if kit.get("logo") and (kit_dir / "logo.png").is_file():
        shutil.copy(kit_dir / "logo.png", build_dir / "brand" / "logo.png")
        theme["logo"] = "brand/logo.png"
        theme["logo_alt"] = kit["logo_alt"]
    files = [ff["file"] for r in ("heading", "body") for ff in kit["fonts"][r].get("files", [])]
    if files:
        (build_dir / "brand" / "fonts").mkdir(exist_ok=True)
        for name in set(files):
            if (kit_dir / "fonts" / name).is_file():
                shutil.copy(kit_dir / "fonts" / name, build_dir / "brand" / "fonts" / name)
    (build_dir / "brand.css").write_text(css(kit), encoding="utf-8")
    return theme
