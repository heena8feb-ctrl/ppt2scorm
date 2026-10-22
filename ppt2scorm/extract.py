"""Step 1 - Extract content from a .pptx: titles, bullet text, speaker notes, tables and teaching visuals.

Visuals:
- Pictures (including ones inside groups and picture placeholders) are saved to media/, resized and compressed.
  Decoration is skipped: images repeated on many slides (logos, template art) and tiny icons.
- Slides whose meaning is in a chart, SmartArt or a diagram drawn from shapes have no picture file to extract,
  so the whole slide is rendered to an image with LibreOffice (when it's installed). Without LibreOffice the
  text is still used and a warning explains what was skipped.
"""
import hashlib
import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_INCH = 914400
MAX_SIDE = 1600              # px, longest side of a saved image
ICON_INCHES = 1.0            # pictures shown smaller than this (both sides) are treated as icons
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
CHART_URI = "http://schemas.openxmlformats.org/drawingml/2006/chart"


# ---------- text ----------
def _shape_texts(shape):
    out = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                out.append({"text": t, "level": p.level})
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
        out.append({"table": rows})
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            out.extend(_shape_texts(s))
    return out


# ---------- visuals ----------
def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)


def _alt_text(shape):
    """The picture's alt text, ignoring PowerPoint's useless defaults like "image.png" or "Picture 3"."""
    for el in shape._element.iter():
        if el.tag.endswith("}cNvPr"):
            alt = (el.get("descr") or "").strip()
            if re.fullmatch(r"(?i).*\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)|(picture|image|graphic|grafik|imagen)\s*\d*", alt):
                return ""
            return alt[:300]
    return ""


def _graphic_uri(shape):
    for el in shape._element.iter():
        if el.tag.endswith("}graphicData"):
            return el.get("uri", "")
    return ""


def _picture(shape):
    """Returns the image of a picture shape or picture placeholder, else None."""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
        try:
            return shape.image
        except (AttributeError, ValueError, KeyError):
            return None
    return None


def _inches(v):
    return (v or 0) / EMU_PER_INCH


def _visual_reason(slide):
    """Why this slide's meaning is in a visual we can't extract as a file (or None)."""
    drawn = 0
    for sh in _walk(slide.shapes):
        if sh.__class__.__name__ == "GraphicFrame":
            uri = _graphic_uri(sh)
            if uri == CHART_URI:
                return "chart"
            if uri == DIAGRAM_URI:
                return "SmartArt diagram"
            continue
        if sh.is_placeholder:
            continue
        if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE) or \
                sh.__class__.__name__ == "Connector":
            text = sh.text_frame.text.strip() if sh.has_text_frame else ""
            if len(text.split()) <= 8:   # boxes with a few words and arrows, not text boxes
                drawn += 1
    return "diagram drawn with shapes" if drawn >= 4 else None


def _save_image(blob, ext, media_dir, digest):
    """Resize/compress an image. Returns (file name, width, height) or None if browsers can't show it."""
    from PIL import Image

    ext = ext.lower()
    if ext in ("emf", "wmf"):
        return None  # Windows vector formats: the slide gets rendered instead
    if ext == "svg":
        return None  # SVG can carry scripts; PowerPoint stores a PNG copy of SVG pictures, which is used instead
    import warnings
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)   # > ~89 megapixels: refuse
            im = Image.open(io.BytesIO(blob))
            im.load()
    except Exception:  # noqa: BLE001 - unreadable or oversized image: skip it
        return None
    if ext == "gif" and getattr(im, "is_animated", False):
        return _write(media_dir, f"img_{digest}.gif", blob, *im.size)
    if max(im.size) > MAX_SIDE:
        im.thumbnail((MAX_SIDE, MAX_SIDE))
    has_alpha = im.mode in ("RGBA", "LA", "P") and ("A" in im.getbands() or "transparency" in im.info)
    buf = io.BytesIO()
    if ext in ("jpg", "jpeg") or (not has_alpha and len(blob) > 300_000 and ext not in ("png", "gif")):
        im.convert("RGB").save(buf, "JPEG", quality=84, optimize=True, progressive=True)
        name = f"img_{digest}.jpg"
    else:
        if im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
            im = im.convert("RGBA" if has_alpha else "RGB")
        im.save(buf, "PNG", optimize=True)
        name = f"img_{digest}.png"
        if not has_alpha and buf.tell() > 400_000:  # a photo saved as PNG: JPEG is far smaller
            buf = io.BytesIO()
            im.convert("RGB").save(buf, "JPEG", quality=84, optimize=True, progressive=True)
            name = f"img_{digest}.jpg"
    return _write(media_dir, name, buf.getvalue(), *im.size)


def _write(media_dir, name, data, w, h):
    (Path(media_dir) / name).write_bytes(data)
    return name, w, h


def find_soffice():
    if os.environ.get("PPT2SCORM_SLIDE_RENDERING") == "0":   # e.g. small free hosting plans with 512 MB of memory
        return None
    for cand in ("soffice", "libreoffice"):
        if shutil.which(cand):
            return shutil.which(cand)
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                 "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(cand):
            return cand
    return None


def render_slides(pptx_path, media_dir, numbers, timeout=150, crop_top=None):
    """Render the given slide numbers to media/slide_N.png. Returns {number: (name, w, h)}.
    crop_top: {number: fraction of the slide height to cut from the top (the title)}; blank margins are trimmed.
    Needs LibreOffice, plus PyMuPDF (pip install pymupdf) or poppler's pdftoppm to turn PDF pages into PNGs."""
    soffice = find_soffice()
    if not soffice or not numbers:
        return {}
    from PIL import Image

    out = {}
    with tempfile.TemporaryDirectory() as tmp:
        # a private profile dir avoids clashing with a LibreOffice window the person already has open
        profile = Path(tmp, "profile").as_uri()
        subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, str(pptx_path)], check=True, capture_output=True, timeout=timeout)
        pdf = next(Path(tmp).glob("*.pdf"))
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(pdf)
            for n in numbers:
                if 1 <= n <= doc.page_count:
                    png = Path(tmp, f"p{n}.png")
                    doc[n - 1].get_pixmap(dpi=110).save(png)
        except ImportError:
            if not shutil.which("pdftoppm"):
                return {}
            for n in numbers:
                subprocess.run(["pdftoppm", "-png", "-r", "110", "-f", str(n), "-l", str(n), "-singlefile",
                                str(pdf), str(Path(tmp, f"p{n}"))], check=True, capture_output=True, timeout=60)
        for n in numbers:
            png = Path(tmp, f"p{n}.png")
            if png.exists():
                im = _trim(Image.open(png).convert("RGB"), (crop_top or {}).get(n, 0))
                im.thumbnail((MAX_SIDE, MAX_SIDE))
                name = f"slide_{n}.png"
                im.save(Path(media_dir, name), "PNG", optimize=True)
                out[n] = (name, *im.size)
    return out


MAX_UNPACKED = int(os.environ.get("PPT2SCORM_MAX_UNPACKED_MB", "400")) * 1024 * 1024
MAX_MEMBERS = 10000


class DeckError(ValueError):
    pass


def check_archive(pptx_path):
    """Refuse decks that would unpack to something huge ("zip bombs") before anything opens them.
    Only the zip's table of contents is read, so this is fast even for big files."""
    import zipfile
    try:
        with zipfile.ZipFile(pptx_path) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, OSError):
        raise DeckError("That file couldn't be opened as a PowerPoint (.pptx) file.") from None
    total = 0
    if len(infos) > MAX_MEMBERS:
        raise DeckError("That deck contains too many parts to be a normal presentation.")
    for i in infos:
        total += i.file_size
        # deflate can't shrink ordinary content more than about 1000x; far more means a crafted file
        if i.file_size > 10 * 1024 * 1024 and i.file_size > 200 * max(i.compress_size, 1):
            raise DeckError("That deck looks damaged or crafted to use up memory, so it wasn't opened.")
    if total > MAX_UNPACKED:
        raise DeckError(f"That deck is over {MAX_UNPACKED // (1024 * 1024)} MB once unpacked. "
                        "Compress its pictures (File > Compress Pictures) and try again.")


def _without_decoration(pptx_path, decorative):
    """A temporary copy of the deck without logos and template pictures, so renders show just the content."""
    prs = Presentation(pptx_path)
    holders = [prs.slides, prs.slide_masters] + [m.slide_layouts for m in prs.slide_masters]
    for group in holders:
        for item in group:
            for sh in list(_walk(item.shapes)):
                img = _picture(sh)
                if img is not None and (img.sha1 in decorative or item.__class__.__name__ != "Slide"):
                    sh._element.getparent().remove(sh._element)
    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    prs.save(path)
    return path


def _trim(im, top_fraction=0.0, pad=24):
    """Cut the title band off a rendered slide and trim blank margins, keeping some padding."""
    from PIL import ImageChops, Image
    w, h = im.size
    if 0 < top_fraction < 0.5:
        im = im.crop((0, int(h * top_fraction), w, h))
    bg = Image.new("RGB", im.size, im.getpixel((2, im.size[1] - 3)))
    box = ImageChops.difference(im, bg).convert("L").point(lambda v: 255 if v > 18 else 0).getbbox()
    if box:
        l, t, r, b = box
        im = im.crop((max(0, l - pad), max(0, t - pad), min(im.size[0], r + pad), min(im.size[1], b + pad)))
    return im


def extract(pptx_path, media_dir, render=True):
    """Returns {title, slides:[{index, title, blocks, notes, images:[{src, w, h, alt, kind}]}], warnings:[...]}"""
    media_dir = Path(media_dir)
    media_dir.mkdir(parents=True, exist_ok=True)
    check_archive(pptx_path)
    prs = Presentation(pptx_path)
    slides, pictures = [], []          # pictures: (slide number, shape, image)
    for i, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        blocks = []
        for shape in slide.shapes:
            if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
                continue
            blocks.extend(_shape_texts(shape))
        for shape in _walk(slide.shapes):
            img = _picture(shape)
            if img is not None:
                pictures.append((i, shape, img))
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        t = slide.shapes.title
        title_bottom = (t.top + t.height) / prs.slide_height if t is not None and t.top is not None and t.height else 0
        slides.append({"index": i, "title": title, "blocks": blocks, "notes": notes, "images": [],
                       "_visual": _visual_reason(slide), "_top": title_bottom})

    # decoration filter: the same image on many slides is a logo or template art
    seen_on = {}
    for n, _, img in pictures:
        seen_on.setdefault(img.sha1, set()).add(n)
    many = max(3, (len(slides) + 1) // 2)
    saved, warnings, unshowable = {}, [], 0
    for n, shape, img in pictures:
        if len(seen_on[img.sha1]) >= many:
            continue
        if _inches(shape.width) < ICON_INCHES and _inches(shape.height) < ICON_INCHES:
            continue
        if img.sha1 not in saved:
            saved[img.sha1] = _save_image(img.blob, img.ext, media_dir, img.sha1[:12])
        res = saved[img.sha1]
        if res is None:
            unshowable += 1
            slides[n - 1]["_visual"] = slides[n - 1]["_visual"] or "picture in a format browsers can't show"
            continue
        name, w, h = res
        entry = {"src": f"media/{name}", "w": w, "h": h, "alt": _alt_text(shape), "kind": "picture"}
        if entry["src"] not in [x["src"] for x in slides[n - 1]["images"]]:
            slides[n - 1]["images"].append(entry)

    # whole-slide renders for charts, SmartArt and shape diagrams
    want = [s["index"] for s in slides if s["_visual"]]
    rendered = {}
    if want and render:
        try:
            clean = _without_decoration(pptx_path, {h for h, on in seen_on.items() if len(on) >= many})
            try:
                rendered = render_slides(clean, media_dir, want, crop_top={x['index']: x['_top'] for x in slides})
            finally:
                os.unlink(clean)
        except (subprocess.SubprocessError, OSError, StopIteration) as e:
            warnings.append(f"Couldn't render slides with charts or diagrams ({e.__class__.__name__}).")
    for s in slides:
        reason = s.pop("_visual")
        s.pop("_top", None)
        if reason and s["index"] in rendered:
            name, w, h = rendered[s["index"]]
            s["images"].append({"src": f"media/{name}", "w": w, "h": h, "kind": "slide",
                                "alt": f"Slide {s['index']}: {s['title'] or reason}", "reason": reason})
    missing = [s for s in want if s not in rendered]
    if missing and render:
        how = "install LibreOffice" if not find_soffice() else "install PyMuPDF (pip install pymupdf)"
        warnings.append(f"{len(missing)} slide(s) show a chart or diagram that couldn't be turned into an image "
                        f"(slides {', '.join(map(str, missing[:8]))}). To include them, {how} and convert again, "
                        "or add an image in the editor.")
    deck_title = slides[0]["title"] if slides and slides[0]["title"] else Path(pptx_path).stem
    return {"title": deck_title, "slides": slides, "warnings": warnings}


def as_prompt_text(deck):
    """Flatten extracted deck into readable text for the LLM."""
    lines = [f"DECK TITLE: {deck['title']}", ""]
    for s in deck["slides"]:
        lines.append(f"--- SLIDE {s['index']}: {s['title']}")
        for b in s["blocks"]:
            if "text" in b:
                lines.append("  " * b["level"] + "- " + b["text"])
            else:
                lines.append("TABLE: " + " | ".join(" / ".join(r) for r in b["table"]))
        for im in s["images"]:
            what = f"whole slide rendered as an image ({im.get('reason', 'diagram')})" if im["kind"] == "slide" else "picture"
            alt = f', alt text "{im["alt"]}"' if im.get("alt") and im["kind"] != "slide" else ""
            lines.append(f"IMAGE: {im['src']} - {what}, {im['w']}x{im['h']} px{alt}")
        if s["notes"]:
            lines.append(f"SPEAKER NOTES: {s['notes']}")
        lines.append("")
    return "\n".join(lines)


def available_images(deck):
    """All image paths the course may reference."""
    return {im["src"]: im for s in deck["slides"] for im in s["images"]}
